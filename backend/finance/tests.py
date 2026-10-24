from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from docx import Document
from pypdf import PdfWriter

from .finance import estimate_uk_tax
from .models import SavingsGoal, Transaction
from .payroll import estimate_job_payroll, gross_for_period
from .payslips import extract_payslip_text, review_payslip_text


class TaxEstimateTests(TestCase):
    def test_basic_rate_employee_estimate(self):
        result = estimate_uk_tax("30000", "england", "2026-27")
        self.assertEqual(result["personal_allowance"], Decimal("12570.00"))
        self.assertEqual(result["income_tax"], Decimal("3486.00"))
        self.assertEqual(result["national_insurance"], Decimal("1394.40"))
        self.assertEqual(result["take_home"], Decimal("25119.60"))

    def test_personal_allowance_taper_and_additional_rate(self):
        result = estimate_uk_tax("150000", "england", "2025-26")
        self.assertEqual(result["personal_allowance"], Decimal("0.00"))
        self.assertEqual(result["income_tax"], Decimal("53703.00"))
        self.assertEqual(result["national_insurance"], Decimal("5010.60"))

    def test_taxable_band_thresholds_remain_income_based_during_allowance_taper(self):
        result = estimate_uk_tax("110000", "england", "2026-27")
        self.assertEqual(result["personal_allowance"], Decimal("7570.00"))
        self.assertEqual(result["taxable_income"], Decimal("102430.00"))
        self.assertEqual(result["income_tax"], Decimal("33432.00"))

    def test_scottish_tax_bands_change_by_year(self):
        older = estimate_uk_tax("30000", "scotland", "2025-26")
        newer = estimate_uk_tax("30000", "scotland", "2026-27")
        self.assertEqual(older["income_tax"], Decimal("3482.82"))
        self.assertEqual(newer["income_tax"], Decimal("3451.07"))

    def test_invalid_inputs_are_rejected(self):
        for args in ((-1, "england", "2026-27"), (100, "england", "2024-25"), (100, "ireland", "2026-27")):
            with self.subTest(args=args), self.assertRaises(ValueError):
                estimate_uk_tax(*args)


class PayrollEstimatorTests(TestCase):
    def test_hourly_pay_is_converted_to_monthly_gross(self):
        gross = gross_for_period(
            pay_basis="hourly", gross_pay=Decimal("0"), hourly_rate=Decimal("15"),
            hours_per_week=Decimal("20"), frequency="monthly",
        )
        self.assertEqual(gross.quantize(Decimal("0.01")), Decimal("1300.00"))

    def test_monthly_1257l_estimate_includes_gross_tax_and_net(self):
        estimate = estimate_job_payroll(
            gross_pay=Decimal("3000"), frequency="monthly", tax_code="1257L",
            region="england", tax_year="2026-27", pay_date=date(2026, 5, 6),
        )
        self.assertEqual(estimate["income_tax"], Decimal("390.40"))
        self.assertEqual(estimate["national_insurance"], Decimal("156.16"))
        self.assertEqual(estimate["net_pay"], Decimal("2453.44"))

    def test_additional_job_br_uses_no_second_personal_allowance(self):
        estimate = estimate_job_payroll(
            gross_pay=Decimal("1000"), frequency="monthly", tax_code="BR",
            region="england", tax_year="2026-27", pay_date=date(2026, 5, 6),
        )
        self.assertEqual(estimate["income_tax"], Decimal("200.00"))

    def test_weekly_w1_tax_code_uses_one_week_basis(self):
        estimate = estimate_job_payroll(
            gross_pay=Decimal("500"), frequency="weekly", tax_code="1257L W1",
            region="england", tax_year="2026-27", pay_date=date(2026, 5, 15),
        )
        self.assertTrue(estimate["period_only"])
        self.assertEqual(estimate["income_tax"], Decimal("51.60"))

    def test_ytd_tax_variance_uses_previous_payday_totals(self):
        estimate = estimate_job_payroll(
            gross_pay=Decimal("3000"), frequency="monthly", tax_code="1257L",
            region="england", tax_year="2026-27", pay_date=date(2026, 5, 6),
            taxable_pay_ytd=Decimal("0"), tax_deducted_before_payday=Decimal("0"),
            actual_tax_this_pay=Decimal("500"),
        )
        self.assertEqual(estimate["ytd_difference"], Decimal("319.20"))


class PayslipExtractionTests(TestCase):
    def test_docx_text_is_read_in_memory_and_values_are_suggestions(self):
        doc = Document()
        doc.add_paragraph("Tax code 1257L")
        doc.add_paragraph("Gross pay £2,500.00")
        doc.add_paragraph("Income tax £300.00")
        stream = BytesIO()
        doc.save(stream)
        upload = SimpleUploadedFile("payslip.docx", stream.getvalue())
        text = extract_payslip_text(upload)
        result = review_payslip_text(text)
        self.assertEqual(result["suggested_tax_code"], "1257L")
        self.assertEqual(result["suggested_gross_pay"], Decimal("2500.00"))
        self.assertEqual(result["suggested_tax_deducted"], Decimal("300.00"))

    def test_unsupported_file_type_is_rejected(self):
        upload = SimpleUploadedFile("payslip.txt", b"1257L")
        with self.assertRaisesMessage(ValueError, "Choose a PDF or Word DOCX"):
            extract_payslip_text(upload)

    def test_encrypted_or_empty_pdf_is_rejected_without_saving(self):
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        stream = BytesIO()
        writer.write(stream)
        upload = SimpleUploadedFile("scan.pdf", stream.getvalue())
        with self.assertRaisesMessage(ValueError, "No selectable text"):
            extract_payslip_text(upload)
        self.assertIsNone(upload.field_name)


class FinanceViewsTests(TestCase):
    def setUp(self):
        self.alice = User.objects.create_user(username="alice", password="CorrectHorse123!")
        self.bob = User.objects.create_user(username="bob", password="CorrectHorse123!")
        self.transaction = Transaction.objects.create(
            user=self.alice, transaction_type=Transaction.EXPENSE,
            amount=Decimal("25.50"), date=date(2026, 5, 3), category="Groceries",
        )

    def test_finance_pages_require_login(self):
        for name in ("dashboard", "transactions", "savings", "tax", "upload"):
            response = self.client.get(reverse(name))
            self.assertEqual(response.status_code, 302)
            self.assertIn(reverse("login"), response["Location"])

    def test_tax_page_renders_four_job_forms(self):
        self.client.force_login(self.alice)
        response = self.client.get(reverse("tax"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'name="job-0-tax_code"')
        self.assertContains(response, 'data-job-index="3"')
        self.assertContains(response, "Extract details to review")

    def test_two_job_form_submission_returns_combined_payday_estimate(self):
        self.client.force_login(self.alice)
        payload = {
            "calculate": "1", "tax_year": "2026-27", "region": "england",
            "pay_date": "2026-05-06", "jobs_count": "2",
            "job-TOTAL_FORMS": "2", "job-INITIAL_FORMS": "0",
            "job-MIN_NUM_FORMS": "0", "job-MAX_NUM_FORMS": "2",
        }
        for index, job_type in enumerate(("main", "additional")):
            payload.update({
                f"job-{index}-job_type": job_type,
                f"job-{index}-pay_frequency": "monthly",
                f"job-{index}-pay_basis": "hourly",
                f"job-{index}-hourly_rate": "15",
                f"job-{index}-hours_per_week": "20",
                f"job-{index}-gross_pay": "",
                f"job-{index}-tax_code": "1257L" if index == 0 else "BR",
                f"job-{index}-taxable_pay_ytd": "",
                f"job-{index}-tax_deducted_ytd": "",
                f"job-{index}-actual_tax_this_pay": "",
            })
        response = self.client.post(reverse("tax"), payload)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "ESTIMATED PAYDAY BREAKDOWN")
        self.assertContains(response, "Main job")
        self.assertContains(response, "Additional job")

    def test_dashboard_and_history_are_scoped_to_signed_in_user(self):
        self.client.force_login(self.bob)
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["transaction_count"], 0)
        self.assertNotContains(response, "Groceries")
        response = self.client.get(reverse("transactions"))
        self.assertNotContains(response, "Groceries")

    def test_user_cannot_edit_or_delete_another_users_transaction(self):
        self.client.force_login(self.bob)
        self.assertEqual(self.client.get(reverse("transaction_edit", args=[self.transaction.pk])).status_code, 404)
        self.assertEqual(self.client.post(reverse("transaction_delete", args=[self.transaction.pk])).status_code, 404)
        self.assertTrue(Transaction.objects.filter(pk=self.transaction.pk).exists())

    def test_create_transaction_is_owned_by_current_user(self):
        self.client.force_login(self.alice)
        response = self.client.post(reverse("transaction_create"), {
            "transaction_type": "expense", "amount": "42.50", "date": "2026-06-01",
            "category": "Loan payment", "description": "Monthly EMI", "currency": "GBP",
        })
        self.assertRedirects(response, reverse("transactions"))
        saved = Transaction.objects.get(description="Monthly EMI")
        self.assertEqual(saved.user, self.alice)

    def test_csv_upload_saves_valid_rows_to_current_user(self):
        self.client.force_login(self.alice)
        file = SimpleUploadedFile(
            "statement.csv",
            b"date,type,amount,category,description,currency\n2026-06-02,expense,320,Loan,Car EMI,GBP\n",
            content_type="text/csv",
        )
        response = self.client.post(reverse("upload"), {"statement": file})
        self.assertRedirects(response, reverse("transactions"))
        self.assertTrue(Transaction.objects.filter(user=self.alice, description="Car EMI", amount="320.00").exists())

    def test_invalid_csv_does_not_import_any_rows(self):
        self.client.force_login(self.alice)
        file = SimpleUploadedFile(
            "statement.csv",
            b"date,type,amount,category\n2026-06-02,expense,-12,Loan\n",
            content_type="text/csv",
        )
        self.client.post(reverse("upload"), {"statement": file})
        self.assertEqual(Transaction.objects.filter(user=self.alice).count(), 1)

    def test_savings_goals_are_user_scoped(self):
        goal = SavingsGoal.objects.create(user=self.alice, name="Emergency fund", target_amount=1000)
        self.client.force_login(self.bob)
        self.assertEqual(self.client.get(reverse("savings_edit", args=[goal.pk])).status_code, 404)
        response = self.client.get(reverse("savings"))
        self.assertNotContains(response, "Emergency fund")
